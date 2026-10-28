# -*- coding: utf-8 -*-
"""基于学校模板生成实训报告 DOCX（v2）。

要点：
1. 标题与正文的字体、字号、行距、缩进全部通过自定义段落样式统一管理
   （实训题目 / 实训章标题 / 实训节标题 / 实训小节标题 / 实训正文 / 实训题注 / 实训图注 / 实训代码 等）；
2. 数据表格全部采用三线表（顶线、底线 1.5 磅，表头下线 0.75 磅，无竖线）；
3. 保留学校模板的封面、任务书、成绩评定表，正文注入在"设计报告正文"标记之后。
"""
import re, sys
import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = r"D:\mqy\java\SpringBoot\online_learning_system\docs\实训报告\报告源稿\report_src.txt"
TPL = r"D:\mqy\java\SpringBoot\online_learning_system\docs\实训报告\实训要求\实践课设计报告模板.docx"
OUT = r"D:\mqy\java\SpringBoot\online_learning_system\docs\实训报告\在线学习平台系统设计与实现-实训报告-莫秋裕.docx"

# ---------- 样式定义 ----------
# (样式名, 段落样式参数, 字体参数)
STYLE_DEFS = [
    ("实训题目",     dict(align=WD_ALIGN_PARAGRAPH.CENTER, space_before=6, space_after=18, line=1.5),
     dict(east="黑体", ascii_f="Times New Roman", size=16, bold=True)),
    ("实训章标题",   dict(align=WD_ALIGN_PARAGRAPH.CENTER, space_before=16, space_after=12, line=1.5,
                          page_break_before=True, keep_with_next=True),
     dict(east="黑体", ascii_f="Times New Roman", size=16, bold=True)),
    ("实训节标题",   dict(align=WD_ALIGN_PARAGRAPH.LEFT, space_before=10, space_after=6, line=1.5,
                          keep_with_next=True),
     dict(east="黑体", ascii_f="Times New Roman", size=14, bold=True)),
    ("实训小节标题", dict(align=WD_ALIGN_PARAGRAPH.LEFT, space_before=8, space_after=4, line=1.5,
                          keep_with_next=True),
     dict(east="黑体", ascii_f="Times New Roman", size=12, bold=True)),
    ("实训正文",     dict(align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent_chars=2, space_before=0, space_after=0, line=1.5),
     dict(east="宋体", ascii_f="Times New Roman", size=12, bold=False)),
    ("实训表格",     dict(align=WD_ALIGN_PARAGRAPH.CENTER, space_before=1, space_after=1, line=1.15),
     dict(east="宋体", ascii_f="Times New Roman", size=10.5, bold=False)),
    ("实训题注",     dict(align=WD_ALIGN_PARAGRAPH.CENTER, space_before=8, space_after=4, line=1.3,
                          keep_with_next=True),
     dict(east="宋体", ascii_f="Times New Roman", size=10.5, bold=True)),
    ("实训图注",     dict(align=WD_ALIGN_PARAGRAPH.CENTER, space_before=2, space_after=10, line=1.3),
     dict(east="宋体", ascii_f="Times New Roman", size=10.5, bold=True)),
    ("实训图占位",   dict(align=WD_ALIGN_PARAGRAPH.CENTER, space_before=8, space_after=2, line=1.3),
     dict(east="楷体", ascii_f="Times New Roman", size=10.5, bold=False, color=(128, 128, 128))),
    ("实训占位说明", dict(align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=4, line=1.3),
     dict(east="楷体", ascii_f="Times New Roman", size=10.5, bold=False, color=(89, 89, 89))),
    ("实训代码",     dict(align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=0, line=1.15),
     dict(east="宋体", ascii_f="Consolas", size=9.5, bold=False)),
    ("实训代码尾注", dict(align=WD_ALIGN_PARAGRAPH.RIGHT, space_before=2, space_after=6, line=1.15),
     dict(east="楷体", ascii_f="Consolas", size=9.5, bold=False, color=(128, 128, 128))),
]

def set_style_font(style, east, ascii_f, size, bold, color=None):
    """在样式级别设置中西文字体、字号、加粗与颜色。"""
    style.font.name = ascii_f
    style.font.size = Pt(size)
    style.font.bold = bold
    if color:
        style.font.color.rgb = RGBColor(*color)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts'); rpr.append(rfonts)
    rfonts.set(qn('w:ascii'), ascii_f)
    rfonts.set(qn('w:hAnsi'), ascii_f)
    rfonts.set(qn('w:eastAsia'), east)

def set_style_indent_chars(style, n):
    """在样式级别设置首行缩进（字符）。"""
    ppr = style.element.get_or_add_pPr()
    ind = ppr.find(qn('w:ind'))
    if ind is None:
        ind = OxmlElement('w:ind'); ppr.append(ind)
    ind.set(qn('w:firstLineChars'), str(n * 100))

def ensure_styles(doc):
    existing = {s.name for s in doc.styles}
    made = {}
    for name, pf, ft in STYLE_DEFS:
        if name in existing:
            st = doc.styles[name]
        else:
            st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            st.quick_style = True
        if pf.get('align') is not None:
            st.paragraph_format.alignment = pf['align']
        st.paragraph_format.space_before = Pt(pf.get('space_before', 0))
        st.paragraph_format.space_after = Pt(pf.get('space_after', 0))
        st.paragraph_format.line_spacing = pf.get('line', 1.5)
        if pf.get('page_break_before'):
            st.paragraph_format.page_break_before = True
        if pf.get('keep_with_next'):
            st.paragraph_format.keep_with_next = True
        if pf.get('indent_chars'):
            set_style_indent_chars(st, pf['indent_chars'])
        set_style_font(st, ft['east'], ft['ascii_f'], ft['size'], ft['bold'], ft.get('color'))
        made[name] = st
    return made

# ---------- 通用小工具 ----------
def insert_paragraph_after(element, doc):
    new_p = OxmlElement('w:p')
    element.addnext(new_p)
    from docx.text.paragraph import Paragraph
    return Paragraph(new_p, doc)

def add_placeholder_border(p):
    pPr = p._element.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    for side in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:val'), 'dashed'); el.set(qn('w:sz'), '8')
        el.set(qn('w:space'), '4'); el.set(qn('w:color'), '999999')
        pbdr.append(el)
    pPr.append(pbdr)

def set_three_line(tbl):
    """三线表：上下框线 1.5 磅，表头下框线 0.75 磅，无竖线与内部横线。"""
    tblPr = tbl._element.tblPr
    borders = OxmlElement('w:tblBorders')
    for side, val, sz in (('top', 'single', 12), ('bottom', 'single', 12),
                          ('left', 'none', 0), ('right', 'none', 0),
                          ('insideH', 'none', 0), ('insideV', 'none', 0)):
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:val'), val); el.set(qn('w:sz'), str(sz))
        el.set(qn('w:space'), '0'); el.set(qn('w:color'), '000000')
        borders.append(el)
    tblPr.append(borders)
    # 表头行下框线 0.75 磅
    for cell in tbl.rows[0].cells:
        tcpr = cell._element.get_or_add_tcPr()
        tcb = OxmlElement('w:tcBorders')
        bt = OxmlElement('w:bottom')
        bt.set(qn('w:val'), 'single'); bt.set(qn('w:sz'), '6')
        bt.set(qn('w:space'), '0'); bt.set(qn('w:color'), '000000')
        tcb.append(bt)
        tcpr.append(tcb)

# ---------- 读取模板 ----------
doc = docx.Document(TPL)
styles = ensure_styles(doc)
body = doc.element.body
from docx.text.paragraph import Paragraph

children = list(body)
mark_body, eval_title = None, None
for child in children:
    if child.tag.endswith('}p'):
        t = Paragraph(child, doc).text.strip()
        if t == '设计报告正文' and mark_body is None:
            mark_body = child
        if t == '本科生实践课程成绩评定表':
            eval_title = child
assert mark_body is not None and eval_title is not None

# 删除"设计报告正文"与"成绩评定表"之间的占位段落
deleting, to_delete = False, []
for child in list(body):
    if child is mark_body:
        deleting = True; continue
    if child is eval_title:
        deleting = False
    if deleting:
        to_delete.append(child)
for el in to_delete:
    body.remove(el)
print(f"已删除占位段落 {len(to_delete)} 个")

cursor = mark_body
count_text = []

def emit(style_name, text=""):
    """以指定样式在游标后插入段落；run 不携带直接字体格式（由样式统一管理）。"""
    global cursor
    p = insert_paragraph_after(cursor, doc)
    p.style = styles[style_name]
    if text:
        p.add_run(text)
    cursor = p._element
    return p

def emit_table(rows):
    """三线表：表注样式已在前，表格单元格段落应用"实训表格"样式，表头加粗。"""
    global cursor
    tbl = doc.add_table(rows=len(rows), cols=len(rows[0]))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = True
    set_three_line(tbl)
    for ri, row in enumerate(rows):
        for ci, txt in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.text = ''
            p = cell.paragraphs[0]
            p.style = styles['实训表格']
            r = p.add_run(txt)
            if ri == 0:
                r.bold = True
    cursor.addnext(tbl._element)
    gap = OxmlElement('w:p')
    tbl._element.addnext(gap)
    gp = Paragraph(gap, doc)
    gp.paragraph_format.space_after = Pt(6)
    cursor = gap

# ---------- 解析源稿并注入 ----------
with open(SRC, encoding='utf-8') as f:
    lines = [l.rstrip('\n') for l in f]

i = 0
while i < len(lines):
    line = lines[i]; s = line.strip()
    i += 1
    if not s:
        continue
    if s.startswith('TITLE|'):
        emit('实训题目', s.split('|', 1)[1])
        count_text.append(s)
    elif s.startswith('H1|'):
        emit('实训章标题', s.split('|', 1)[1])
        count_text.append(s)
    elif s.startswith('H2|'):
        emit('实训节标题', s.split('|', 1)[1])
        count_text.append(s)
    elif s.startswith('H3|'):
        emit('实训小节标题', s.split('|', 1)[1])
        count_text.append(s)
    elif s.startswith('TABLE|'):
        cap = s.split('|', 1)[1]
        emit('实训题注', cap)
        count_text.append(cap)
        rows = []
        while i < len(lines) and lines[i].strip().startswith('R|'):
            rows.append([c.strip() for c in lines[i].strip()[2:].split('|')])
            count_text.append(lines[i].strip()[2:].replace('|', ''))
            i += 1
        emit_table(rows)
    elif s.startswith('FIG|'):
        parts = s.split('|')
        cap, note = parts[1], parts[2] if len(parts) > 2 else ''
        ph = emit('实训图占位', '【图片占位框：导出/截取图片后替换本框，删除本框文字】')
        add_placeholder_border(ph)
        emit('实训占位说明', note)
        emit('实训图注', cap)
        count_text.append(cap + note)
    elif s.startswith('CODE|'):
        emit('实训代码', s.split('|', 1)[1] if s.count('|') else '')
        count_text.append(s)
        while i < len(lines) and lines[i].strip().startswith('CODE|'):
            emit('实训代码', lines[i].strip().split('|', 1)[1])
            count_text.append(lines[i].strip())
            i += 1
        emit('实训代码尾注', '（代码片段结束）')
    elif s.startswith('P|'):
        emit('实训正文', s.split('|', 1)[1])
        count_text.append(s)
    else:
        emit('实训正文', s)
        count_text.append(s)

doc.save(OUT)

full = ''.join(count_text)
cjk = len(re.findall(r'[\u4e00-\u9fff\u3000-\u303f\uff00-\uffef]', full))
nonspace = len(re.sub(r'\s', '', full))
print(f"保存至: {OUT}")
print(f"中文字符+中文标点: {cjk}")
print(f"非空白字符总数(近似Word不计空格字符数): {nonspace}")
